from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import io
import re


def generate_cover_letter_docx(cover_letter_text: str, job_title: str, company: str) -> bytes:
    """
    Generates a formatted .docx cover letter.
    Returns bytes so FastAPI can stream it as a file download.
    """
    doc = Document()

    # Standard professional margins
    for section in doc.sections:
        section.top_margin    = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin   = Inches(1.25)
        section.right_margin  = Inches(1.25)

    # Heading
    title_para = doc.add_heading(level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run("Cover Letter")
    run.font.size      = Pt(16)
    run.font.color.rgb = RGBColor(0x1e, 0x40, 0xaf)  # dark blue

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run(f"{job_title}  ·  {company}")
    sub_run.font.size      = Pt(11)
    sub_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)

    doc.add_paragraph()  # blank line spacer

    # Body — split on double newlines (paragraph breaks)
    paragraphs = [p.strip() for p in cover_letter_text.split("\n\n") if p.strip()]
    for para_text in paragraphs:
        para = doc.add_paragraph()
        run  = para.add_run(para_text)
        run.font.size                       = Pt(11)
        para.paragraph_format.space_after   = Pt(10)
        para.paragraph_format.line_spacing  = Pt(14)

    # Save to in-memory bytes buffer
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


_KNOWN_HEADERS = {
    'PROFESSIONAL SUMMARY', 'SUMMARY', 'PROFESSIONAL PROFILE', 'OBJECTIVE',
    'EDUCATION', 'ACADEMIC BACKGROUND',
    'CERTIFICATIONS', 'CERTIFICATION', 'LICENSES & CERTIFICATIONS',
    'PROFESSIONAL EXPERIENCE', 'WORK EXPERIENCE', 'EXPERIENCE', 'EMPLOYMENT HISTORY',
    'PROJECTS', 'KEY PROJECTS', 'SELECTED PROJECTS',
    'TECHNICAL SKILLS', 'SKILLS', 'CORE COMPETENCIES', 'KEY SKILLS',
    'ACHIEVEMENTS', 'ACCOMPLISHMENTS', 'AWARDS', 'HONORS',
    'PUBLICATIONS', 'LANGUAGES', 'VOLUNTEER EXPERIENCE',
}


def _add_bottom_border(para):
    try:
        pPr = para._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '1e40af')
        pBdr.append(bottom)
        pPr.append(pBdr)
    except Exception:
        pass  # falls back gracefully — header still renders as bold


def _is_contact_line(line):
    return '@' in line or line.count('|') >= 2 or bool(re.search(r'\+?\d[\d\s\-]{8,}', line))


def _is_section_header(line):
    stripped = line.strip()
    upper = stripped.upper()
    if upper in _KNOWN_HEADERS:
        return True
    # Generic: short all-caps alphabetic line (e.g. "LANGUAGES", "AWARDS")
    if upper == stripped and 2 <= len(stripped) <= 45 and stripped.replace(' ', '').replace('&', '').replace('-', '').isalpha():
        return True
    return False


def _is_company_role_line(line):
    stripped = line.strip()
    if not stripped:
        return False
    if stripped[0] in '•–—*▪':
        return False
    if stripped.startswith('- '):
        return False
    # Em dash (—) or en dash (–) typically separates company from role title
    return ' — ' in stripped or ' – ' in stripped


def generate_resume_docx(resume_text: str, job_title: str, company: str) -> bytes:
    """
    Generates an ATS-approved .docx of the AI-rewritten resume.
    Processes line-by-line to properly format name, contact, headers, and bullets.
    Returns bytes for FastAPI file download.
    """
    doc = Document()

    # ATS-safe margins: 1.0" all sides
    for sec in doc.sections:
        sec.top_margin    = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin   = Inches(1.0)
        sec.right_margin  = Inches(1.0)

    lines = resume_text.split('\n')
    non_empty_seen = 0

    for line in lines:
        stripped = line.strip()

        if not stripped:
            continue

        non_empty_seen += 1

        # ── Name (first non-empty line) ──────────────────────────────────────
        if non_empty_seen == 1:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(stripped)
            run.bold           = True
            run.font.size      = Pt(18)
            run.font.color.rgb = RGBColor(0x1e, 0x40, 0xaf)
            run.font.name      = 'Calibri'
            para.paragraph_format.space_after = Pt(2)
            continue

        # ── Contact info (second non-empty line with email/phone/pipes) ──────
        if non_empty_seen == 2 and _is_contact_line(stripped):
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(stripped)
            run.font.size      = Pt(10)
            run.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
            run.font.name      = 'Calibri'
            para.paragraph_format.space_after = Pt(8)
            continue

        # ── Section header ───────────────────────────────────────────────────
        if _is_section_header(stripped):
            para = doc.add_paragraph()
            run = para.add_run(stripped.upper())
            run.bold           = True
            run.font.size      = Pt(11)
            run.font.color.rgb = RGBColor(0x1e, 0x40, 0xaf)
            run.font.name      = 'Calibri'
            para.paragraph_format.space_before = Pt(14)
            para.paragraph_format.space_after  = Pt(3)
            _add_bottom_border(para)
            continue

        # ── Bullet item ──────────────────────────────────────────────────────
        if stripped[0] in '•–—*▪' or stripped.startswith('- '):
            # Normalize to simple • (ATS-safe on Taleo/Workday)
            content = re.sub(r'^[•–—▪\*\-]\s*', '', stripped)
            para = doc.add_paragraph()
            run = para.add_run(f'•  {content}')
            run.font.size = Pt(10.5)
            run.font.name = 'Calibri'
            para.paragraph_format.left_indent        = Inches(0.25)
            para.paragraph_format.first_line_indent  = Inches(-0.18)
            para.paragraph_format.space_after        = Pt(3)
            continue

        # ── Company / role line (e.g. "Deloitte USI — Software Engineer II") ─
        if _is_company_role_line(stripped):
            para = doc.add_paragraph()
            run = para.add_run(stripped)
            run.bold      = True
            run.font.size = Pt(10.5)
            run.font.name = 'Calibri'
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after  = Pt(3)
            continue

        # ── Regular body text ────────────────────────────────────────────────
        para = doc.add_paragraph()
        run = para.add_run(stripped)
        run.font.size = Pt(10.5)
        run.font.name = 'Calibri'
        para.paragraph_format.space_after  = Pt(4)
        para.paragraph_format.line_spacing = Pt(13.5)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()